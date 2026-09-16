#!/usr/bin/env python3
"""Gera o MODELO da cotação comparativa do Grupo Leh.

Este script é FERRAMENTA DE BUILD, não roda a cada pedido. Ele produz o
template que fica no OneDrive; o n8n copia esse template por solicitação e
escreve só o cabeçalho e os itens. As fórmulas moram no template, escritas
uma vez, aqui.

O QUE MUDOU EM RELAÇÃO AO MODELO ANTIGO

  · 3 colunas de fornecedor, não 5 — o card do ClickUp comporta 3, e uma
    planilha que aceita 5 deixaria o 4º existir na planilha e não no card,
    que é o que o aprovador lê.
  · Empresa no cabeçalho — entrou no pedido em 16/09 e é ela que vai no topo
    da ordem de compra do GR.
  · LINHAS DE ITEM: 200, não 60. O modelo antigo calculava até a linha 70, e
    um 61º item era escrito e sumia de todas as contas — com a planilha ainda
    dizendo "Cotou todos os itens" e "0 itens em branco". Nada acusava. É a
    mesma família do erro que elegeu vencedor quem não tinha cotado.
  · A coluna "Referência internet" continua, preenchida à mão pelo comprador
    quando ele cotar na internet.

AS TRÊS TRAVAS, que são o motivo desta planilha existir:

  1. Célula vazia e preço zero NUNCA vencem a disputa de menor preço.
  2. Fornecedor que deixou item em branco não entra na comparação por total —
     ele continua ganhando itens individuais, só não disputa o pedido inteiro.
  3. A economia de comprar item a item já desconta os fretes a mais que a
     divisão do pedido causaria. Só entra o frete de quem ganha alguma coisa.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

ITENS = 200          # linhas de item
LIN1  = 11           # primeira linha de item
FORN  = 3            # colunas de fornecedor

# ---- geometria, derivada, para não existir número mágico solto -------------
ULT      = LIN1 + ITENS - 1          # última linha de item
C_F1     = 8                         # H
C_FN     = C_F1 + FORN - 1           # J
C_REF    = C_FN + 1                  # K  Referência internet
C_MENOR  = C_REF + 1                 # L
C_VENC   = C_MENOR + 1               # M
C_PORKG  = C_VENC + 1                # N
C_TOTAL  = C_PORKG + 1               # O
C_OCULTA = C_TOTAL + 1               # P..  totais por fornecedor
ULT_VIS  = C_TOTAL

R_RESUMO = ULT + 3
R_TOTITE = R_RESUMO + 1
R_CAB    = R_RESUMO + 3
R_F1     = R_CAB + 1
R_FN     = R_F1 + FORN - 1
R_SOMA   = R_FN + 2
R_FRETES = R_SOMA + 1
R_TOTAL  = R_SOMA + 2
R_UNICO  = R_SOMA + 3
R_ECON   = R_SOMA + 4
R_PCT    = R_SOMA + 5
R_NOTA   = R_PCT + 3

AZUL   = '0E2438'                     # o azul da marca, tirado do próprio logo
CINZA  = 'E9EDF2'                     # vem do formulário
AMAREL = 'FFF3C4'                     # o comprador preenche
AZULC  = 'DCEAF7'                     # referência da internet
BORDA  = Side(style='thin', color='C7CED6')
TODAS  = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)

def montar(caminho):
    wb = Workbook(); ws = wb.active; ws.title = 'Comparativo'

    # ---------------- título ----------------
    ws.cell(1,1,'COTAÇÃO COMPARATIVA — GRUPO LEH')
    ws.cell(1,1).font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    ws.cell(1,1).alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ULT_VIS)
    for c in range(1, ULT_VIS+1): ws.cell(1,c).fill = PatternFill('solid', fgColor=AZUL)
    ws.row_dimensions[1].height = 26

    # ---------------- cabeçalho ----------------
    # Esquerda em A(label) / C(valor); direita em F(label) / H(valor).
    cab = [(3,'Nº da solicitação:','Centro de custo:'),
           (4,'Solicitante:','Tipo de compra:'),
           (5,'Local de entrega:','Entrega até:'),
           (6,'Motivo:','Empresa:')]
    for r, esq, dir_ in cab:
        ws.cell(r,1,esq); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.cell(r,6,dir_); ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        for c in (1,6):
            ws.cell(r,c).font = Font(name='Arial', size=10, bold=True)
            ws.cell(r,c).alignment = Alignment(horizontal='right')
        for c in (3,8):
            ws.cell(r,c).font = Font(name='Arial', size=10)
            ws.cell(r,c).fill = PatternFill('solid', fgColor=CINZA)
            ws.cell(r,c).border = TODAS
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=ULT_VIS)

    # ---------------- legenda ----------------
    ws.cell(8,1,'DE ONDE VEM CADA COISA   ·   CINZA: do formulário de solicitação   ·   '
                'AMARELO: você preenche com as propostas dos fornecedores   ·   '
                'AZUL: referência de preço que você pesquisou na internet   ·   '
                'o resto é calculado pela planilha')
    ws.cell(8,1).font = Font(name='Arial', size=9, italic=True, color='45505E')
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=ULT_VIS)
    ws.row_dimensions[8].height = 22

    # ---------------- títulos das colunas ----------------
    titulos = ['#','Código','Item','Qtd','Unid.','Conteúdo por unidade','Medida (kg ou L)']
    titulos += ['Fornecedor %d' % (i+1) for i in range(FORN)]
    titulos += ['Referência internet','Menor preço','Vencedor','Preço por kg/L','Total do item']
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(10, i, t)
        c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = TODAS
    ws.row_dimensions[10].height = 34
    for i in range(FORN):
        ws.cell(10, C_F1+i).fill = PatternFill('solid', fgColor='2A4A66')
    ws.cell(10, C_REF).fill = PatternFill('solid', fgColor='2F5C85')

    larguras = {1:5, 2:12, 3:44, 4:8, 5:11, 6:15, 7:14, C_REF:14,
                C_MENOR:12, C_VENC:15, C_PORKG:14, C_TOTAL:14}
    for i in range(FORN): larguras[C_F1+i] = 14
    for c, w in larguras.items(): ws.column_dimensions[L(c)].width = w

    # ---------------- linhas de item ----------------
    fx = lambda c, r: '$%s%d' % (L(c), r)
    pF = lambda r: '$%s%d:$%s%d' % (L(C_F1), r, L(C_FN), r)     # H..J da linha
    for r in range(LIN1, ULT+1):
        ws.cell(r,1, '=IF(%s="","",ROW()-%d)' % (fx(3,r), LIN1-1))
        # TRAVA 1 — vazio e zero nunca vencem. SMALL ignora branco e texto;
        # o COUNTIF pula quantos forem <= 0.
        ws.cell(r, C_MENOR,
            '=IF(%s="","",IFERROR(SMALL(%s,COUNTIF(%s,"<=0")+1),""))' % (fx(3,r), pF(r), pF(r)))
        ws.cell(r, C_VENC,
            '=IF(%s="","",IFERROR(INDEX($%s$10:$%s$10,MATCH(%s,%s,0)),""))'
            % (fx(C_MENOR,r), L(C_F1), L(C_FN), fx(C_MENOR,r), pF(r)))
        ws.cell(r, C_PORKG,
            '=IF(OR(%s="",NOT(ISNUMBER(%s)),%s<=0),"",%s/%s)'
            % (fx(C_MENOR,r), fx(6,r), fx(6,r), fx(C_MENOR,r), fx(6,r)))
        ws.cell(r, C_TOTAL,
            '=IF(OR(%s="",NOT(ISNUMBER(%s))),"",%s*%s)'
            % (fx(C_MENOR,r), fx(4,r), fx(C_MENOR,r), fx(4,r)))
        # colunas ocultas: total daquele fornecedor naquele item
        for i in range(FORN):
            ws.cell(r, C_OCULTA+i,
                '=IF(OR(%s="",NOT(ISNUMBER(%s)),%s<=0,NOT(ISNUMBER(%s))),"",%s*%s)'
                % (fx(3,r), fx(C_F1+i,r), fx(C_F1+i,r), fx(4,r), fx(4,r), fx(C_F1+i,r)))

        for c in range(1, ULT_VIS+1):
            cel = ws.cell(r,c)
            cel.border = TODAS
            cel.font = Font(name='Arial', size=10)
            if c in (2,3,4,5,6,7): cel.fill = PatternFill('solid', fgColor=CINZA)
            elif C_F1 <= c <= C_FN: cel.fill = PatternFill('solid', fgColor=AMAREL)
            elif c == C_REF:        cel.fill = PatternFill('solid', fgColor=AZULC)
        for c in list(range(C_F1, C_REF+1)) + [C_MENOR, C_PORKG, C_TOTAL]:
            ws.cell(r,c).number_format = 'R$ #,##0.00'
        ws.cell(r,4).number_format = '0.###'
    for i in range(FORN):
        ws.column_dimensions[L(C_OCULTA+i)].hidden = True

    # ---------------- resumo ----------------
    ws.cell(R_RESUMO,1,'RESUMO DA COTAÇÃO')
    ws.cell(R_RESUMO,1).font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    ws.merge_cells(start_row=R_RESUMO, start_column=1, end_row=R_RESUMO, end_column=ULT_VIS)
    for c in range(1, ULT_VIS+1): ws.cell(R_RESUMO,c).fill = PatternFill('solid', fgColor=AZUL)

    ws.cell(R_TOTITE,3,'Total de itens na solicitação:')
    ws.cell(R_TOTITE,6,'=COUNTA($C$%d:$C$%d)' % (LIN1, ULT))

    for i, t in enumerate(['Fornecedor','Itens cotados','Itens em branco','Subtotal dos itens',
                           'Custo de entrega','Total comparável','Situação'], start=3):
        c = ws.cell(R_CAB, i, t)
        c.font = Font(name='Arial', size=10, bold=True)
        c.fill = PatternFill('solid', fgColor='D8E0E8'); c.border = TODAS
    ws.merge_cells(start_row=R_CAB, start_column=9, end_row=R_CAB, end_column=ULT_VIS)

    for i in range(FORN):
        r  = R_F1 + i
        oc = L(C_OCULTA+i)                      # coluna oculta deste fornecedor
        faixa = '$%s$%d:$%s$%d' % (oc, LIN1, oc, ULT)
        ws.cell(r,3,'=IF(COUNT(%s)=0,"",$%s$10)' % (faixa, L(C_F1+i)))
        ws.cell(r,4,'=IF($C%d="","",COUNT(%s))' % (r, faixa))
        ws.cell(r,5,'=IF($C%d="","",$F$%d-$D%d)' % (r, R_TOTITE, r))
        # TRAVA 2 — quem deixou item em branco não soma total comparável
        ws.cell(r,6,'=IF(OR($C%d="",$E%d>0),"",SUM(%s))' % (r, r, faixa))
        ws.cell(r,8,'=IF($F%d="","",$F%d+N($G%d))' % (r, r, r))
        ws.cell(r,9,'=IF($C%d="","",IF($E%d>0,"Deixou "&$E%d&" item(ns) em branco — fora da '
                    'comparação por total","Cotou todos os itens"))' % (r, r, r))
        # TRAVA 3 — só entra o frete de quem ganha alguma coisa
        ws.cell(r, C_OCULTA,
            '=IF($C%d="",0,IF(COUNTIF($%s$%d:$%s$%d,$C%d)>0,N($G%d),0))'
            % (r, L(C_VENC), LIN1, L(C_VENC), ULT, r, r))
        for c in range(3, ULT_VIS+1):
            ws.cell(r,c).border = TODAS; ws.cell(r,c).font = Font(name='Arial', size=10)
        ws.cell(r,7).fill = PatternFill('solid', fgColor=AMAREL)   # frete: ele preenche
        for c in (6,7,8): ws.cell(r,c).number_format = 'R$ #,##0.00'
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=ULT_VIS)

    linhas = [
      (R_SOMA,  'Soma dos itens comprando cada um do mais barato:',
                '=SUM($%s$%d:$%s$%d)' % (L(C_TOTAL), LIN1, L(C_TOTAL), ULT)),
      (R_FRETES,'Fretes dos fornecedores que seriam acionados:',
                '=SUM($%s$%d:$%s$%d)' % (L(C_OCULTA), R_F1, L(C_OCULTA), R_FN)),
      (R_TOTAL, 'TOTAL comprando item a item:', '=$F$%d+$F$%d' % (R_SOMA, R_FRETES)),
      (R_UNICO, 'Melhor total de um fornecedor único (com frete):',
                '=IFERROR(SMALL($H$%d:$H$%d,COUNTIF($H$%d:$H$%d,"<=0")+1),"")'
                % (R_F1, R_FN, R_F1, R_FN)),
      (R_ECON,  'Economia comprando item a item:',
                '=IF(OR($F$%d="",$F$%d=0),"",$F$%d-$F$%d)' % (R_UNICO, R_TOTAL, R_UNICO, R_TOTAL)),
      (R_PCT,   'Economia (%):',
                '=IF(OR($F$%d="",$F$%d=0),"",$F$%d/$F$%d)' % (R_UNICO, R_UNICO, R_ECON, R_UNICO)),
    ]
    for r, rot, f in linhas:
        ws.cell(r,3,rot); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws.cell(r,3).font = Font(name='Arial', size=10, bold=(r in (R_TOTAL, R_ECON)))
        ws.cell(r,3).alignment = Alignment(horizontal='right')
        cel = ws.cell(r,6,f); cel.border = TODAS
        cel.font = Font(name='Arial', size=10, bold=(r in (R_TOTAL, R_ECON)))
        cel.number_format = '0.0%' if r == R_PCT else 'R$ #,##0.00'

    ws.cell(R_NOTA,3,
      'Três travas propositais: célula vazia e preço zero nunca vencem a disputa de menor '
      'preço; fornecedor que deixou item em branco não entra na comparação por total (mas '
      'continua ganhando itens individuais); e a economia de comprar item a item já desconta '
      'os fretes a mais que a divisão do pedido causaria.')
    ws.cell(R_NOTA,3).font = Font(name='Arial', size=9, italic=True, color='45505E')
    ws.cell(R_NOTA,3).alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=R_NOTA, start_column=3, end_row=R_NOTA+2, end_column=ULT_VIS)

    ws.freeze_panes = 'A11'
    wb.save(caminho)
    return dict(itens=ITENS, primeira=LIN1, ultima=ULT, fornecedores=FORN,
                col_f1=L(C_F1), col_fn=L(C_FN), col_ref=L(C_REF),
                linha_res=R_F1, col_frete='G')

if __name__ == '__main__':
    destino = sys.argv[1] if len(sys.argv) > 1 else 'cotacao/MODELO-Cotacao.xlsx'
    info = montar(destino)
    print(destino)
    for k, v in info.items(): print('  %-14s %s' % (k, v))
