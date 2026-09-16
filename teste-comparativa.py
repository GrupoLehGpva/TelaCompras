#!/usr/bin/env python3
"""Prova das fórmulas do MODELO-Cotacao gerado.

Não confere se a planilha abre: confere se as CONTAS estão certas, com um
cenário montado para ser conferível à mão e recalculado pelo LibreOffice.

Recalcular sem erro só prova que as fórmulas avaliam. Um intervalo errado por
uma linha dá um arquivo limpo com número errado — e número errado numa cotação
sai na fatura, não na tela.
"""
import openpyxl, shutil, subprocess, os, sys

AQUI    = os.path.dirname(os.path.abspath(__file__))
MODELO  = os.path.join(AQUI, 'cotacao', 'MODELO-Cotacao.xlsx')
TRAB    = '/tmp/claude-0/-home-claude/9ea1e790-347f-553b-96ac-7d693ea21849/scratchpad/cot'
os.makedirs(TRAB, exist_ok=True)

LIN1, ULT = 11, 210
R_TOT, R_F1 = 214, 217
R_SOMA, R_FRETES, R_TOTAL, R_UNICO, R_ECON, R_PCT = 221, 222, 223, 224, 225, 226

falhas = []
def confere(nome, obtido, esperado, tol=0.005):
    ok = (abs(obtido-esperado) < tol) if isinstance(esperado,(int,float)) and isinstance(obtido,(int,float)) \
         else obtido == esperado
    print(('ok      ' if ok else 'FALHOU  ') + nome.ljust(54) + ' = ' + repr(obtido) +
          ('' if ok else '   esperado ' + repr(esperado)))
    if not ok: falhas.append(nome)

def montar(nome, itens, precos, fretes):
    p = os.path.join(TRAB, nome); shutil.copy(MODELO, p)
    wb = openpyxl.load_workbook(p); ws = wb['Comparativo']
    for i,(c,d,f,g) in enumerate(itens):
        r = LIN1+i
        ws.cell(r,3).value=c; ws.cell(r,4).value=d
        ws.cell(r,6).value=f; ws.cell(r,7).value=g
    for r,cols in precos.items():
        for col,v in cols.items(): ws.cell(r,col).value = v
    for i,v in enumerate(fretes): ws.cell(R_F1+i,7).value = v
    wb.save(p); return p

def recalcular(p):
    saida = os.path.join(TRAB,'calc'); os.makedirs(saida, exist_ok=True)
    subprocess.run(['soffice','--headless','--norestore','--convert-to',
                    'xlsx:Calc MS Excel 2007 XML','--outdir',saida,p],
                   check=True, capture_output=True, timeout=240)
    return os.path.join(saida, os.path.basename(p))

def ler(p, cels):
    ws = openpyxl.load_workbook(p, data_only=True)['Comparativo']
    return {c: ws[c].value for c in cels}

# =====================================================================
# CENARIO A — as tres travas de uma vez, com 3 fornecedores
#   item 1 SABONETE   10un  0,09 kg   F1 2,00 · F2 1,80 · F3 ZERO
#   item 2 DETERGENTE  5un  0,5  L    F1 3,00 · F2 3,50 · F3 2,90
#   item 3 LUVA       20un  sem cont. F1 1,00 · F2 1,00 (EMPATE) · F3 5,00
#   fretes  F1 50 · F2 30 · F3 100
# =====================================================================
p = recalcular(montar('A.xlsx',
  [('SABONETE',10,0.09,'kg'),('DETERGENTE',5,0.5,'L'),('LUVA',20,None,None)],
  {11:{8:2.00, 9:1.80, 10:0},
   12:{8:3.00, 9:3.50, 10:2.90},
   13:{8:1.00, 9:1.00, 10:5.00}},
  [50,30,100]))
v = ler(p, ['L11','M11','N11','O11','L12','M12','O12','M13','N13','O13',
            'F214','C217','D217','E217','F217','H217','I217',
            'D218','F218','H218','C219','D219','E219','F219','H219','I219',
            'F221','F222','F223','F224','F225','F226'])

print('--- TRAVA 1: vazio e zero nunca vencem ---')
confere('item 1 menor preco ignora o zero do F3',      v['L11'], 1.80)
confere('item 1 vencedor',                             v['M11'], 'Fornecedor 2')
confere('item 1 preco por kg = 1,80 / 0,09',           v['N11'], 20.00)
confere('item 1 total = 1,80 x 10',                    v['O11'], 18.00)
confere('item 2 menor = 2,90 (F3)',                    v['L12'], 2.90)
confere('item 2 vencedor',                             v['M12'], 'Fornecedor 3')
confere('item 3 sem conteudo: preco por kg vazio',     v['N13'], None)
confere('empate fica com o primeiro da esquerda',      v['M13'], 'Fornecedor 1')

print('\n--- TRAVA 2: quem deixou branco sai da comparacao por total ---')
confere('total de itens',                              v['F214'], 3)
confere('F1 cotou 3',                                  v['D217'], 3)
confere('F1 subtotal 20+15+20',                        v['F217'], 55.00)
confere('F1 total comparavel 55 + 50',                 v['H217'], 105.00)
confere('F1 situacao',                                 v['I217'], 'Cotou todos os itens')
confere('F2 subtotal 18+17,5+20',                      v['F218'], 55.50)
confere('F2 total comparavel 55,5 + 30',               v['H218'], 85.50)
confere('F3 cotou so 2',                               v['D219'], 2)
confere('F3 deixou 1 em branco',                       v['E219'], 1)
confere('F3 FORA da comparacao por total',             v['F219'], None)
confere('F3 sem total comparavel',                     v['H219'], None)
confere('F3 situacao explica', v['I219'], 'Deixou 1 item(ns) em branco — fora da comparação por total')

print('\n--- TRAVA 3: so o frete de quem ganha alguma coisa ---')
confere('soma item a item 18 + 14,5 + 20',             v['F221'], 52.50)
confere('fretes acionados: os tres ganharam algo',     v['F222'], 180.00)
confere('TOTAL item a item',                           v['F223'], 232.50)
confere('melhor fornecedor unico = F2',                v['F224'], 85.50)
confere('economia NEGATIVA (o frete come a diferenca)',v['F225'], -147.00)
confere('economia %',                                  v['F226'], -147.00/85.50)

# =====================================================================
# CENARIO B — o limite. 200 itens tem que fechar a conta inteira.
# =====================================================================
p = recalcular(montar('B.xlsx',
  [('ITEM %03d'%i,1,None,None) for i in range(1,201)],
  {LIN1+i:{8:10.0} for i in range(200)}, [0,0,0]))
v = ler(p, ['F214','D217','E217','F217','I217','F221','A210','C210','L210'])
print('\n--- LIMITE: 200 itens, o teto do modelo ---')
confere('a linha 210 e item, e tem numero',            v['A210'], 200)
confere('e calcula menor preco',                       v['L210'], 10.0)
confere('total de itens = 200',                        v['F214'], 200)
confere('fornecedor cotou os 200',                     v['D217'], 200)
confere('nenhum em branco',                            v['E217'], 0)
confere('subtotal = 200 x 10',                         v['F217'], 2000.0)
confere('soma item a item = 2000',                     v['F221'], 2000.0)

print('\n===== FALHAS (%d) =====' % len(falhas))
for f in falhas: print(' x ' + f)
sys.exit(1 if falhas else 0)
