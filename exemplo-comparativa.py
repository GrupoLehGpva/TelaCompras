#!/usr/bin/env python3
"""Gera um EXEMPLO preenchido, para conferir de olho antes de ligar no fluxo.

O modelo em si nasce vazio de propósito: uma linha de exemplo dentro dele
viraria um item fantasma na primeira cotação de verdade.
"""
import openpyxl, shutil, subprocess, os
AQUI = os.path.dirname(os.path.abspath(__file__))
MOD  = os.path.join(AQUI,'cotacao','MODELO-Cotacao.xlsx')
OUT  = os.path.join(AQUI,'cotacao','EXEMPLO-C2609-00001-Comparativo.xlsx')
shutil.copy(MOD, OUT)
wb = openpyxl.load_workbook(OUT); ws = wb['Comparativo']

# cabeçalho — é o que o n8n vai escrever, vindo da solicitação
ws['C3'] = 'C2609-00001'
ws['H3'] = '3180 · MANUTENÇÃO'
ws['C4'] = 'MARIA EDUARDA ANTUNES MACIEL SIQUEIRA'
ws['H4'] = 'Normal'
ws['C5'] = 'Fazenda Noricum — Guarapuava/PR'
ws['H5'] = '30/09/2026'
ws['C6'] = 'Reposição do estoque de limpeza da casa dos moradores'
ws['H6'] = 'WIENFRIED MATTHIAS LEH - PR'

# itens — também do n8n
itens = [
  ('8210','BIOPLUS 2B (20 KG)',            200, 'KG',       20,  'kg'),
  ('LIM-04','SABONETE LÍQUIDO 5L',          12, 'Galão',      5,  'L'),
  ('LIM-11','DETERGENTE NEUTRO 500ML',      48, 'Frasco',   0.5,  'L'),
  ('EPI-02','LUVA NITRÍLICA TAM. G',        20, 'Par',     None, None),
]
for i,(cod,desc,qtd,un,cont,med) in enumerate(itens):
    r = 11+i
    ws.cell(r,2,cod); ws.cell(r,3,desc); ws.cell(r,4,qtd)
    ws.cell(r,5,un);  ws.cell(r,6,cont); ws.cell(r,7,med)

# o que o HERISSON preenche (amarelo) — três propostas, uma incompleta
precos = {11:{8:38.807, 9:39.50, 10:38.90},
          12:{8:92.00,  9:88.40, 10:None},
          13:{8:4.15,   9:4.15,  10:3.98},
          14:{8:6.80,   9:6.20,  10:6.20}}
for r,cols in precos.items():
    for c,v in cols.items():
        if v is not None: ws.cell(r,c,v)
ws['K11'] = 41.90            # referência da internet, que ele pesquisou
ws['G217'], ws['G218'], ws['G219'] = 180.00, 120.00, 260.00   # fretes
wb.save(OUT)

saida = os.path.dirname(OUT)
subprocess.run(['soffice','--headless','--norestore','--convert-to',
                'xlsx:Calc MS Excel 2007 XML','--outdir','/tmp/exrecalc',OUT],
               check=True, capture_output=True, timeout=240)
shutil.copy('/tmp/exrecalc/' + os.path.basename(OUT), OUT)

v = openpyxl.load_workbook(OUT, data_only=True)['Comparativo']
print('EXEMPLO gerado:', OUT)
for cel,rot in [('F214','total de itens'),('F221','soma item a item'),
                ('F222','fretes acionados'),('F223','TOTAL item a item'),
                ('F224','melhor fornecedor único'),('F225','economia'),('F226','economia %')]:
    print('  %-24s %s' % (rot, v[cel].value))
for r in (217,218,219):
    print('  F%d: %s | cotados %s | branco %s | %s' %
          (r-216, v['C%d'%r].value, v['D%d'%r].value, v['E%d'%r].value, v['I%d'%r].value))
